"""Unit tests: reading files into pages."""

import io
import shutil
import zipfile

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from course_assistant.ingest import (
    IngestError,
    drop_repeated_lines,
    file_hash,
    _find_soffice,
    prepare_pptx,
    unlink_text,
    read_document,
    read_text,
)

# The same lookup the app uses, so a Mac install in /Applications counts too
HAS_LIBREOFFICE = shutil.which(_find_soffice()) is not None


def test_pdf_pages_keep_text_and_images(sample_pdf, tmp_path):
    pages = read_document(sample_pdf, tmp_path / "work")
    assert [p.number for p in pages] == [1, 2, 3]
    assert pages[0].label == "page"
    assert "retrieves relevant information" in pages[0].text
    for page in pages:
        assert page.image_path and (tmp_path / "work" / "pages").exists()


def test_markdown_splits_into_sections(sample_md):
    pages = read_text(sample_md)
    assert [p.title for p in pages] == ["Office hours", "Grading"]
    assert all(p.label == "section" and not p.image_path for p in pages)


def test_unsupported_type_is_rejected(tmp_path):
    path = tmp_path / "data.xlsx"
    path.write_bytes(b"not a real spreadsheet")
    with pytest.raises(IngestError, match="isn't a supported file type"):
        read_document(path, tmp_path / "work")


def test_empty_text_file_is_rejected(tmp_path):
    path = tmp_path / "empty.txt"
    path.write_text("   \n", encoding="utf-8")
    with pytest.raises(IngestError, match="empty"):
        read_document(path, tmp_path / "work")


def test_repeated_headers_and_footers_are_removed():
    texts = [f"Topic {chr(64 + i)}\nreal content {i * 10}\nCourse Syllabus (Subject to Change) {i}" for i in range(1, 7)]
    cleaned = drop_repeated_lines(texts)
    assert all("Subject to Change" not in t for t in cleaned)
    assert cleaned[2].startswith("Topic C")
    assert "real content 30" in cleaned[2]


def test_same_bytes_give_same_hash(sample_pdf, tmp_path):
    copy = tmp_path / "renamed.pdf"
    shutil.copyfile(sample_pdf, copy)
    assert file_hash(sample_pdf) == file_hash(copy)


@pytest.mark.skipif(not HAS_LIBREOFFICE, reason="LibreOffice is not installed")
def test_powerpoint_is_converted_to_slide_images(tmp_path):
    deck = Presentation()
    for title in ["Prompt engineering", "Structured output"]:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = f"Notes about {title.lower()}"
    path = tmp_path / "week.pptx"
    deck.save(path)
    pages = read_document(path, tmp_path / "work")
    assert [p.label for p in pages] == ["slide", "slide"]
    assert pages[1].title == "Structured output"
    assert all(p.image_path for p in pages)


def _deck_with_animated_gif(tmp_path):
    # Starts blank, then draws a navy box: the box is the frame worth keeping
    frames = [Image.new("RGB", (64, 48), "white") for _ in range(3)]
    for frame in frames[1:]:
        frame.paste("navy", (8, 8, 56, 40))
    gif = io.BytesIO()
    frames[0].save(gif, format="GIF", save_all=True, append_images=frames[1:], duration=100, loop=0)
    gif.seek(0)
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Terminal demo"
    slide.shapes.add_picture(gif, Inches(1), Inches(2), width=Inches(6))
    path = tmp_path / "gif-deck.pptx"
    deck.save(path)
    return path


def test_animated_gifs_become_their_most_detailed_frame(tmp_path):
    deck = _deck_with_animated_gif(tmp_path)
    frozen_path, frozen = prepare_pptx(deck, tmp_path / "prepared")
    assert frozen == 1
    with zipfile.ZipFile(deck) as before, zipfile.ZipFile(frozen_path) as after:
        assert before.namelist() == after.namelist()  # same parts, so slide links still work
        gif_name = next(n for n in after.namelist() if n.endswith(".gif"))
        with Image.open(io.BytesIO(after.read(gif_name))) as still:
            assert getattr(still, "n_frames", 1) == 1
            red, green, blue = still.convert("RGB").getpixel((32, 24))
            assert blue > 100 and red < 50  # navy: the drawn frame, not the blank first one


def test_non_zip_file_is_left_alone(tmp_path):
    path = tmp_path / "old.pptx"
    path.write_bytes(b"not a zip")
    assert prepare_pptx(path, tmp_path / "prepared") == (path, 0)


@pytest.mark.skipif(not HAS_LIBREOFFICE, reason="LibreOffice is not installed")
def test_slide_with_animated_gif_shows_the_drawn_frame(tmp_path):
    # End-to-end check: the converted slide shows the navy box, not the blank first frame.
    pages = read_document(_deck_with_animated_gif(tmp_path), tmp_path / "work")
    with Image.open(pages[0].image_path) as image:
        navy = sum(1 for r, g, b in image.convert("RGB").getdata() if b > 100 and r < 50 and g < 50)
    assert navy > 1000  # the drawn frame, not an empty box


def test_unreadable_gif_is_left_for_libreoffice(tmp_path):
    deck = _deck_with_animated_gif(tmp_path)
    broken = tmp_path / "broken.pptx"
    with zipfile.ZipFile(deck) as source, zipfile.ZipFile(broken, "w") as target:
        for item in source.infolist():
            data = source.read(item)
            target.writestr(item, b"GIF89a not really an image" if item.filename.endswith(".gif") else data)
    frozen_path, frozen = prepare_pptx(broken, tmp_path / "prepared")
    assert frozen == 0 and zipfile.ZipFile(frozen_path).testzip() is None


@pytest.mark.skipif(not HAS_LIBREOFFICE, reason="LibreOffice is not installed")
def test_prepared_copy_is_removed_after_conversion(tmp_path):
    read_document(_deck_with_animated_gif(tmp_path), tmp_path / "work")
    assert not (tmp_path / "work" / "prepared").exists()


NS = (
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
)


def test_linked_text_becomes_plain_text_styled_as_a_link():
    slide = (
        f'<p:txBody xmlns:p="p" {NS}><a:p>'
        '<a:r><a:rPr lang="en-US"/><a:t>Plain </a:t></a:r>'  # self-closing settings must stay untouched
        '<a:r><a:rPr lang="en-US" dirty="0"><a:hlinkClick r:id="rId10"/></a:rPr>'
        "<a:t>Artificial Analysis Intelligence Index</a:t></a:r>"
        '<a:r><a:rPr lang="en-US"><a:solidFill><a:srgbClr val="FF0000"/></a:solidFill>'
        '<a:hlinkClick r:id="rId11"/></a:rPr><a:t>Red link</a:t></a:r>'
        "</a:p></p:txBody>"
    ).encode()
    changed, count = unlink_text(slide)
    assert count == 2 and b"hlinkClick" not in changed
    assert b'<a:r><a:rPr lang="en-US"/><a:t>Plain </a:t></a:r>' in changed
    assert b'<a:rPr lang="en-US" dirty="0" u="sng"><a:solidFill><a:schemeClr val="hlink"/></a:solidFill></a:rPr>' in changed
    assert b'<a:srgbClr val="FF0000"/>' in changed and changed.count(b"schemeClr") == 1  # its own colour is kept


def test_unlinking_never_returns_a_broken_slide():
    broken = f'<a:p {NS}><a:r><a:rPr><a:hlinkClick r:id="rId1"/></a:rPr><a:t>x</a:t></a:r>'.encode()  # unclosed
    assert unlink_text(broken) == (broken, 0)


@pytest.mark.skipif(not HAS_LIBREOFFICE, reason="LibreOffice is not installed")
def test_long_link_wraps_between_words(tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(2.2), Inches(2)).text_frame
    box.word_wrap = True
    run = box.paragraphs[0].add_run()
    run.text = "Artificial Analysis Intelligence Index"
    run.hyperlink.address = "https://artificialanalysis.ai"
    path = tmp_path / "links.pptx"
    deck.save(path)
    text = read_document(path, tmp_path / "work")[0].text
    assert "Intelligence" in text.split()  # not "I" + "ntelligence"
